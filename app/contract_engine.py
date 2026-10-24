# -*- coding: utf-8 -*-
"""合同识别引擎：本地文本提取 + 规则识别，预留可配置的大模型适配器

设计原则：
- 默认完全本地运行，不依赖外部大模型 API；
- 规则引擎（rules）：正则 + 中文大写金额转换，确定性提取金额/税率/付款方式/节点比例；
- 大模型适配器（llm）：可通过环境变量配置 OpenAI 兼容接口，配置后优先使用，失败自动降级规则引擎；
- 文本提取：PDF（pdfplumber）/ Word（python-docx）/ TXT-MD（编码探测）/ 图片（Windows 自带 OCR）。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import uuid

# ---------- OCR 文本清理 ----------

def _clean_ocr_text(text):
    """清理 Windows OCR 常见噪声：移除中文字符间空格、数字间空格、全角转半角、修复常见误识别"""
    if not text:
        return text
    # 1. 将全角数字/字母/标点转为半角
    trans = str.maketrans(
        '０１２３４５６７８９　，．／；：！？＠＃＄％＾＆＊（）＿＋｛｝［］｜＼＜＞',
        '0123456789 ,./;:!?@#$%^&*()_+{}[]|\\<>'
    )
    text = text.translate(trans)
    # 2. 移除中文字符之间的空格
    text = re.sub(r'([\u4e00-\u9fff])\s+([\u4e00-\u9fff])', r'\1\2', text)
    text = re.sub(r'([\u4e00-\u9fff])\s+([\u4e00-\u9fff])', r'\1\2', text)
    # 3. 移除数字与逗号/点之间的空格（如 165 , 000 . 00 -> 165,000.00）
    text = re.sub(r'([0-9])\s*[,，]\s*([0-9])', r'\1,\2', text)
    text = re.sub(r'([0-9])\s*\.\s*([0-9])', r'\1.\2', text)
    text = re.sub(r'([0-9])\s+([0-9])', r'\1\2', text)
    # 4. 移除货币符号与数字之间的空格
    text = re.sub(r'([￥¥$€])\s+([0-9])', r'\1\2', text)
    # 5. 修复常见 OCR 误识别
    text = text.replace('圆', '元')
    return text


# ---------- 中文大写金额转换 ----------
CN_NUM = {'零': 0, '壹': 1, '贰': 2, '叁': 3, '肆': 4, '伍': 5,
          '陆': 6, '柒': 7, '捌': 8, '玖': 9,
          '〇': 0, '一': 1, '二': 2, '三': 3, '四': 4,
          '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}
CN_UNIT = {'拾': 10, '佰': 100, '仟': 1000,
           '十': 10, '百': 100, '千': 1000}
CN_SECTION = {'万': 10000, '亿': 100000000}


def _parse_section(sec):
    total = 0
    cur = 0
    for ch in sec:
        if ch in CN_NUM:
            cur = CN_NUM[ch]
        elif ch in CN_UNIT:
            if cur == 0:
                cur = 1
            total += cur * CN_UNIT[ch]
            cur = 0
        else:
            if cur:
                total += cur
                cur = 0
    if cur:
        total += cur
    return total


def chinese_amount_to_number(s):
    if not s:
        return None
    s = (s.replace('人民币', '').replace('元整', '').replace('圆整', '')
         .replace('元', '').replace('圆', '').replace('整', '').replace('正', '')
         .replace('角', '').replace('分', ''))
    if not s:
        return None
    if '亿' in s:
        left, _, right = s.partition('亿')
        return _parse_section(left or '') * 100000000 + _parse_section(right or '')
    if '万' in s:
        left, _, right = s.partition('万')
        return _parse_section(left or '') * 10000 + _parse_section(right or '')
    return _parse_section(s)


AMOUNT_KEYWORDS = r'(?:合同(?:总)?(?:金额|价款|价)|合同总价|总金额|合同金额|项目金额|成交金额|采购金额)'
AMOUNT_RE_DIGIT = re.compile(
    r'[^。；;\n]{0,36}' + AMOUNT_KEYWORDS + r'[^。；;\n]{0,36}')
AMOUNT_NUM = re.compile(r'([0-9][0-9,，]*(?:\.[0-9]{1,2})?)')
AMOUNT_CN = re.compile(r'([零〇壹贰叁肆伍陆柒捌玖一二三四五六七八九十拾佰仟万亿圆元角分整正]{4,})')
AMOUNT_SYMBOL = re.compile(r'[￥¥]\s*([0-9][0-9,，]*(?:\.[0-9]{1,2})?)')
AMOUNT_BARE = re.compile(r'\b([1-9][0-9]{4,}(?:\.[0-9]{1,2})?)\b')


def extract_amount(text):
    """提取合同金额，返回 (金额, 出处片段) 或 (None, None)"""
    for m in AMOUNT_RE_DIGIT.finditer(text):
        seg = m.group(0).strip()
        for n in AMOUNT_NUM.findall(seg):
            try:
                v = float(n.replace(',', '').replace('，', ''))
            except ValueError:
                continue
            if 500 <= v <= 999999999:
                return v, seg
        for c in AMOUNT_CN.findall(seg):
            v = chinese_amount_to_number(c)
            if v and 500 <= v <= 999999999:
                return v, seg
    m = AMOUNT_SYMBOL.search(text)
    if m:
        try:
            v = float(m.group(1).replace(',', '').replace('，', ''))
            if 500 <= v <= 999999999:
                return v, m.group(0).strip()
        except ValueError:
            pass
    m = AMOUNT_BARE.search(text)
    if m:
        try:
            v = float(m.group(1))
            if 500 <= v <= 999999999:
                return v, m.group(0).strip()
        except ValueError:
            pass
    return None, None


def extract_tax_rate(text):
    """提取税率，返回 (税率字符串, 出处片段)"""
    m1 = re.search(r'(?:增值税)?税率\s*[:：]?\s*(?:为)?\s*([0-9]{1,2}(?:\.[0-9])?)\s*%', text)
    if m1:
        return m1.group(1) + '%', m1.group(0).strip()
    m2 = re.search(r'[^。；;\n]{0,24}(?:增值税|税率)[^。；;\n]{0,10}?([0-9]{1,2}(?:\.[0-9])?)\s*%', text)
    if m2:
        return m2.group(1) + '%', m2.group(0).strip()
    m3 = re.search(r'([0-9]{1,2}(?:\.[0-9])?)\s*%\s*(?:的增值税|增值税|税率)', text)
    if m3:
        return m3.group(1) + '%', m3.group(0).strip()
    return None, None


# 各类节点比例提取（首个命中为准）
NODE_PATTERNS = {
    'first': [
        re.compile(r'(?:首付款?|预付款?|预付定?金|定金|签约款)[^。；;\n]{0,24}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
        re.compile(r'(?:合同签订|签订合同|合同生效)[^。；;\n]{0,20}?(?:后|之日起)?[^。；;\n]{0,10}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
    ],
    'progress': [
        re.compile(r'进度款?[^。；;\n]{0,24}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
        re.compile(r'(?:货到|到货|交付|安装)[^。；;\n]{0,20}?(?:后)?[^。；;\n]{0,12}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
    ],
    'acceptance': [
        re.compile(r'(?:验收款?|尾款|完工结算|余额|余款)[^。；;\n]{0,24}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
    ],
    'warranty': [
        re.compile(r'(?:质保金|保修金|质保款|保证金)[^。；;\n]{0,24}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%'),
        re.compile(r'[^。；;\n]{0,24}?([0-9]{1,3}(?:\.[0-9]{1,2})?)\s*%\s*(?:作?为(?:质保金|保证金|保修金)|质保金|保证金)'),
    ],
}

#: CR-04：系统不再提供任何默认比例，未识别比例留空由人工填写。
NODE_NAMES = {'first': '首付款', 'progress': '进度款', 'acceptance': '验收款', 'warranty': '质保金'}
CALC_BASES = {'first': 'total', 'progress': 'total', 'acceptance': 'total', 'warranty': 'total'}


def _rule_node_name(node_type, ratio):
    """节点名称：类型 + 实际比例；比例未识别时只写类型名。"""
    from .nodes import node_name as _node_name
    return _node_name(node_type, ratio)


def extract_ratios(text):
    """提取各节点付款比例，返回 {node_type: (ratio浮点|None, 出处片段, 是否来自合同原文)}"""
    result = {}
    for nt, pats in NODE_PATTERNS.items():
        found = False
        for p in pats:
            m = p.search(text)
            if m:
                try:
                    r = float(m.group(1)) / 100.0
                except (ValueError, KeyError):
                    continue
                result[nt] = (r, m.group(0).strip(), True)
                found = True
                break
        if not found:
            result[nt] = (None, None, False)
    return result


def extract_payment_method(text):
    """提取付款方式 / 付款条件，返回 (摘要, 出处片段)"""
    for kw in ('付款方式', '支付方式', '结算方式'):
        m = re.search(re.escape(kw) + r'\s*[:：]?\s*([^。；;\n]{2,60})', text)
        if m and not re.match(r'^[，。、]', m.group(1)):
            return m.group(1).strip(), (m.group(0).strip()[:160])
    for m in re.finditer(r'[^。；;\n]{0,80}(?:分期支付|一次性付清|按月支付|按进度支付|预付)[^。；;\n]{0,40}', text):
        seg = m.group(0).strip()
        if seg:
            return seg[:80], seg[:160]
    return None, None


# ---------- 文本提取 ----------

def _detect_encoding(path):
    with open(path, 'rb') as f:
        raw = f.read(64 * 1024)
    for enc in ('utf-8-sig', 'utf-8', 'gb18030'):
        try:
            raw.decode(enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return 'gb18030'


def extract_text(path, file_type):
    """按文件类型提取全文，图片走 Windows 自带 OCR；返回 (text, 提示信息)"""
    if file_type == 'pdf':
        return _extract_pdf(path)
    if file_type == 'docx':
        return _extract_docx(path)
    if file_type in ('txt', 'md'):
        return _extract_txt(path)
    if file_type == 'img':
        return _extract_image(path)
    return None, '不支持的文件类型'


def _extract_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        return None, '缺少 PDF 解析库 pdfplumber（pip install pdfplumber）'
    try:
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ''
                if t.strip():
                    parts.append(t)
        text = '\n'.join(parts)
        if text.strip():
            return text, None
        # pdfplumber 未提取到文字，尝试用 PyMuPDF 渲染为图片后 OCR
        return _extract_scanned_pdf(path)
    except Exception as e:
        return None, 'PDF 解析失败：%s' % e


# ---------- Windows OCR（WinRT 只能读系统临时目录，先转存再识别） ----------

_OCR_STAGING_SUBDIR = 'payment_ocr'


def _ocr_staging_dir():
    """OCR 暂存目录（WinRT 可访问、且进程可写）。

    工作区/instance 路径在 WinRT ``StorageFile`` 下会报 0x80070005（拒绝访问）；
    OCR 前统一把图片复制到这里，识别完立即清理。优先系统 TEMP（沙箱/普通
    进程都可写），其次 %LOCALAPPDATA%\\Temp\\payment_ocr。
    """
    import tempfile
    candidates = [os.path.join(tempfile.gettempdir(), _OCR_STAGING_SUBDIR)]
    local = os.environ.get('LOCALAPPDATA')
    if local:
        candidates.append(os.path.join(local, 'Temp', _OCR_STAGING_SUBDIR))
    last_error = None
    for path in candidates:
        try:
            os.makedirs(path, exist_ok=True)
            return path
        except OSError as exc:
            last_error = exc
    raise OSError('无法创建 OCR 暂存目录：%s' % last_error)


def _stage_image(source_path):
    """把图片复制到 OCR 暂存目录（文件名唯一），返回暂存路径。"""
    suffix = os.path.splitext(source_path)[1].lower() or '.png'
    staged = os.path.join(_ocr_staging_dir(), 'ocr_%s%s' % (uuid.uuid4().hex, suffix))
    shutil.copyfile(source_path, staged)
    return staged


def _remove_quietly(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except OSError:
        pass


#: OCR 失败时用于区分「文件访问被拒」的标记（PS 脚本与 .NET 报错文本）
_ACCESS_DENIED_MARKS = ('0x80070005', 'Access is denied', '拒绝访问',
                        '__OCR_ACCESS_DENIED__')


def _run_ocr(staged_path):
    """对 WinRT 可访问的图片路径执行 OCR；返回 (text, error)。"""
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    script = os.path.join(base, 'scripts', 'ocr_win.ps1')
    if not os.path.exists(script):
        return None, '未找到本地 OCR 脚本 ocr_win.ps1'
    try:
        proc = subprocess.run(
            ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass',
             '-File', script, '-ImagePath', staged_path],
            capture_output=True, text=True, encoding='utf-8',
            errors='replace', timeout=120)
    except subprocess.TimeoutExpired:
        return None, '图片 OCR 超时'
    except Exception as e:
        return None, '图片 OCR 失败：%s' % e

    out = (proc.stdout or '').strip()
    combined = out + '\n' + ((proc.stderr or '').strip())
    if '__NO_OCR_ENGINE__' in combined:
        return None, ('系统未安装中文 OCR 语言包，无法识别图片合同；'
                      '请上传文字版 PDF / Word / TXT')
    if any(mark in combined for mark in _ACCESS_DENIED_MARKS):
        return None, ('图片 OCR 无法读取文件（访问被拒 0x80070005）：'
                      '图片路径对 Windows OCR 不可访问，请重试或上传文字版合同')
    if proc.returncode != 0:
        detail = (combined.strip().splitlines() or ['未知错误'])[-1]
        return None, '图片 OCR 失败：%s' % detail[:300]
    if not out:
        return None, '图片 OCR 未能识别到文字，请上传清晰图片或文字版合同'
    return _clean_ocr_text(out), None


def _extract_scanned_pdf(path):
    """扫描件 PDF：用 PyMuPDF 渲染页面为图片，再调用 Windows OCR 识别"""
    try:
        import fitz
    except ImportError:
        return None, '扫描件 PDF 识别需要 PyMuPDF（pip install PyMuPDF）'
    try:
        doc = fitz.open(path)
        if doc.page_count == 0:
            return None, 'PDF 为空文件'
        ocr_parts = []
        for i in range(doc.page_count):
            page = doc.load_page(i)
            # 使用 200 DPI 保证清晰度
            mat = fitz.Matrix(200 / 72, 200 / 72)
            pix = page.get_pixmap(matrix=mat)
            # 直接渲染到 WinRT 可访问的暂存目录，识别后立即删除
            img_path = os.path.join(_ocr_staging_dir(),
                                    'scan_%s_%d.png' % (uuid.uuid4().hex, i))
            pix.save(img_path)
            t, err = _run_ocr(img_path)
            if t:
                ocr_parts.append(t)
            _remove_quietly(img_path)
        doc.close()
        if not ocr_parts:
            return None, '扫描件 PDF 未能识别到文字，请上传清晰文件或文字版合同'
        return _clean_ocr_text('\n'.join(ocr_parts)), None
    except Exception as e:
        return None, '扫描件 PDF 识别失败：%s' % e


def _extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        return None, '缺少 Word 解析库 python-docx（pip install python-docx）'
    try:
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(' | '.join(cells))
        text = '\n'.join(parts)
        if not text.strip():
            return None, 'Word 文档内容为空或为旧版 .doc，请另存为 .docx 后上传'
        return text, None
    except Exception as e:
        return None, 'Word 解析失败：%s' % e


def _extract_txt(path):
    try:
        enc = _detect_encoding(path)
        with open(path, 'r', encoding=enc, errors='replace') as f:
            text = f.read()
        return text, None
    except Exception as e:
        return None, '文本读取失败：%s' % e


def _extract_image(path):
    """调用 Windows 自带 OCR（scripts/ocr_win.ps1），全部本地处理。

    先复制到 ``%LOCALAPPDATA%\\Temp\\payment_ocr`` 再识别：WinRT ``StorageFile``
    访问工作区/instance 路径会报 0x80070005（拒绝访问），这就是
    “同一张图放系统临时目录能识别”的根因。
    """
    if not os.path.exists(path):
        return None, '图片文件不存在或已被清理'
    staged = None
    try:
        staged = _stage_image(path)
        return _run_ocr(staged)
    except OSError as e:
        return None, '图片 OCR 暂存失败（无法复制到系统临时目录）：%s' % e
    finally:
        _remove_quietly(staged)


# ---------- 大模型适配器（可配置） ----------

def load_llm_config():
    """读取大模型识别配置（环境变量），未配置时返回 None（纯本地模式）。

    配置项：
      CONTRACT_LLM_ENDPOINT   OpenAI 兼容接口地址（如 https://dashscope.aliyuncs.com/compatible-mode/v1）
      CONTRACT_LLM_API_KEY    API Key
      CONTRACT_LLM_MODEL      模型名（如 qwen-plus / deepseek-chat）
      CONTRACT_LLM_TIMEOUT    超时秒数（默认 60）
    """
    endpoint = os.environ.get('CONTRACT_LLM_ENDPOINT', '').strip()
    api_key = os.environ.get('CONTRACT_LLM_API_KEY', '').strip()
    model = os.environ.get('CONTRACT_LLM_MODEL', '').strip()
    if not (endpoint and api_key and model):
        return None
    return {
        'endpoint': endpoint.rstrip('/'),
        'api_key': api_key,
        'model': model,
        'timeout': int(os.environ.get('CONTRACT_LLM_TIMEOUT', '60') or 60),
    }


def recognize_with_llm(text, file_name=''):
    """调用配置好的大模型识别，返回结构化结果；失败返回 None 由规则引擎兜底。"""
    cfg = load_llm_config()
    if not cfg:
        return None, '未配置大模型接口'
    try:
        import urllib.request
        prompt = (
            '你是合同信息抽取助手。请从下面的合同文本中提取字段，只输出 JSON，不要输出其他内容。\n'
            '输出格式：\n'
            '{\n'
            '  "contract_amount": 165000,\n'
            '  "tax_rate": "13%",\n'
            '  "payment_method": "分期支付",\n'
            '  "nodes": [{"node_type": "first", "node_name": "首付款30%", "ratio": 0.3}],\n'
            '  "source_texts": {"contract_amount": "原文片段", "tax_rate": "原文片段",'
            ' "payment_method": "原文片段", "first": "原文片段"}\n'
            '}\n'
            '节点类型仅用 first(首付/预付)/progress(进度款)/acceptance(验收款)/warranty(质保金)。\n'
            '金额保留原始数值（合同总金额），比例按百分比转小数；原文没有明确比例时 ratio 输出 null，'
            '不要推测或补默认值。\n\n合同文本：\n' + text[:8000]
        )
        payload = json.dumps({
            'model': cfg['model'],
            'messages': [{'role': 'user', 'content': prompt}],
            'temperature': 0.1,
        }).encode('utf-8')
        req = urllib.request.Request(
            cfg['endpoint'] + '/chat/completions', data=payload,
            headers={'Content-Type': 'application/json',
                     'Authorization': 'Bearer ' + cfg['api_key']})
        with urllib.request.urlopen(req, timeout=cfg['timeout']) as resp:
            data = json.loads(resp.read().decode('utf-8'))
        content = data['choices'][0]['message']['content']
        content = content.strip()
        if content.startswith('```'):
            content = re.sub(r'^```(?:json)?\s*', '', content)
            content = re.sub(r'\s*```$', '', content)
        obj = json.loads(content)
        return obj, None
    except Exception as e:
        return None, '大模型识别失败（%s），已降级为本地规则识别' % e


# ---------- 规则识别主入口 ----------

def recognize_by_rules(text, file_name=''):
    """规则引擎识别，返回与 LLM 一致的结构化结果"""
    amount, amount_src = extract_amount(text)
    tax_rate, tax_src = extract_tax_rate(text)
    pay_method, pay_src = extract_payment_method(text)
    ratios = extract_ratios(text)

    nodes = []
    total = amount or 0.0
    for nt in ('first', 'progress', 'acceptance', 'warranty'):
        ratio, src, from_text = ratios[nt]
        base = total if CALC_BASES[nt] == 'total' else (amount or 0.0)
        due = round(base * ratio, 2) if ratio is not None else None
        node_name = _rule_node_name(nt, ratio)
        nodes.append({
            'node_type': nt,
            'node_name': node_name,
            'ratio': ratio,
            'calc_base': CALC_BASES[nt],
            'amount': due,
            'source_text': src or '',
            'from_text': from_text,
        })

    sources = {'contract_amount': amount_src or '', 'tax_rate': tax_src or '',
               'payment_method': pay_src or ''}
    for nd in nodes:
        if nd['source_text']:
            sources.setdefault(nd['node_type'], nd['source_text'])

    warnings = []
    if amount is None:
        warnings.append('未识别到合同金额，请手动填写')
    if not amount_src:
        warnings.append('合同金额未定位到原文出处，请核对')
    if tax_rate is None:
        warnings.append('未识别到税率，请核对合同')
    un_ratio = [NODE_NAMES[k] for k in ('first', 'progress', 'acceptance', 'warranty')
                if not ratios[k][2]]
    if un_ratio:
        warnings.append('未识别到%s比例，请在确认页按合同人工填写（系统不提供默认比例）'
                        % '/'.join(un_ratio))

    return {
        'engine': 'rules',
        'contract_amount': amount,
        'tax_rate': tax_rate,
        'payment_method': pay_method,
        'nodes': nodes,
        'sources': sources,
        'warnings': warnings,
    }


def recognize_contract(text, file_name=''):
    """识别主入口：优先配置的大模型，失败/未配置用本地规则引擎"""
    llm_result, llm_msg = recognize_with_llm(text, file_name)
    if llm_result:
        try:
            nodes = []
            for nd in llm_result.get('nodes', []):
                nt = nd.get('node_type')
                if nt not in ('first', 'progress', 'acceptance', 'warranty'):
                    continue
                raw_ratio = nd.get('ratio')
                ratio = None if raw_ratio in (None, '') else float(raw_ratio)
                total = float(llm_result.get('contract_amount') or 0)
                base = total if CALC_BASES.get(nt) == 'total' else total
                nodes.append({
                    'node_type': nt,
                    'node_name': nd.get('node_name') or _rule_node_name(nt, ratio),
                    'ratio': ratio,
                    'calc_base': CALC_BASES.get(nt, 'total'),
                    'amount': round(base * ratio, 2) if ratio is not None else None,
                    'source_text': (llm_result.get('source_texts') or {}).get(nt, ''),
                    'from_text': ratio is not None,
                })
            return {
                'engine': 'llm',
                'contract_amount': float(llm_result.get('contract_amount') or 0) or None,
                'tax_rate': llm_result.get('tax_rate'),
                'payment_method': llm_result.get('payment_method'),
                'nodes': nodes,
                'sources': llm_result.get('source_texts') or {},
                'warnings': [llm_msg] if llm_msg else [],
            }, None
        except Exception as e:
            pass  # 结构异常降级规则
    result = recognize_by_rules(text, file_name)
    if llm_msg:
        result['warnings'].append(llm_msg)
    return result, None


def _default_nodes(total):
    """兼容名称保留：不再生成默认比例节点（CR-04 删除全部默认值）。"""
    return []
